#!/usr/bin/env python3
"""
AWVS 자동 엑셀 보고서 생성 모듈
스캔 결과 JSON → 4개 시트 .xlsx 보고서 자동 생성

사용법:
    from generate_report import generate_xlsx_report
    filepath = generate_xlsx_report(scan_data, attack_mapping)
"""

import os
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── 공통 스타일 ──────────────────────────────────────────
FONT_HEADER = Font(name="Arial", bold=True, size=11, color="FFFFFF")
FONT_TITLE = Font(name="Arial", bold=True, size=16)
FONT_SUBTITLE = Font(name="Arial", bold=True, size=11)
FONT_LABEL = Font(name="Arial", bold=True, size=10)
FONT_NORMAL = Font(name="Arial", size=10)
FONT_SMALL = Font(name="Arial", size=9, color="555555")

FILL_HEADER_BLUE = PatternFill("solid", fgColor="2B5797")
FILL_HEADER_ORANGE = PatternFill("solid", fgColor="E8833A")
FILL_HEADER_PURPLE = PatternFill("solid", fgColor="7C5CFC")
FILL_HEADER_RED = PatternFill("solid", fgColor="E05252")

FILL_VULN = PatternFill("solid", fgColor="FFCCCC")
FILL_WARN = PatternFill("solid", fgColor="FFF3CD")
FILL_SAFE = PatternFill("solid", fgColor="D4EDDA")
FILL_NA = PatternFill("solid", fgColor="E2E3E5")
FILL_LIGHT_GRAY = PatternFill("solid", fgColor="F2F2F2")
FILL_DARK_HEADER = PatternFill("solid", fgColor="2B3E50")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_WRAP = Alignment(vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

STATUS_FILL_MAP = {
    "취약": FILL_VULN,
    "주의": FILL_WARN,
    "양호": FILL_SAFE,
    "N/A": FILL_NA,
}

STATUS_FONT_MAP = {
    "취약": Font(name="Arial", bold=True, size=10, color="CC0000"),
    "주의": Font(name="Arial", bold=True, size=10, color="856404"),
    "양호": Font(name="Arial", bold=True, size=10, color="155724"),
    "N/A": Font(name="Arial", size=10, color="666666"),
}

RISK_COLORS = {
    "Critical": "CC0000", "High": "E05252",
    "Medium": "E8833A", "Low": "28A745",
}


def _apply_header_row(ws, row, headers, fill, widths=None):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = FONT_HEADER
        cell.fill = fill
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    if widths:
        for col_idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w


def _apply_cell(ws, row, col, value, font=None, fill=None, alignment=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font or FONT_NORMAL
    cell.border = THIN_BORDER
    cell.alignment = alignment or ALIGN_LEFT
    if fill:
        cell.fill = fill
    return cell


def _section_header(ws, row, col_start, col_end, title, fill):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=title)
    cell.font = FONT_SUBTITLE
    cell.fill = fill
    cell.alignment = ALIGN_LEFT
    cell.border = THIN_BORDER
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).fill = fill
        ws.cell(row=row, column=c).border = THIN_BORDER


# ── Sheet 1: 스캔 요약 ──────────────────────────────────
def _build_summary_sheet(wb, scan_data, results):
    ws = wb.active
    ws.title = "스캔 요약"

    vuln_count = sum(1 for r in results if r["status"] == "취약")
    warn_count = sum(1 for r in results if r["status"] == "주의")
    safe_count = sum(1 for r in results if r["status"] == "양호")
    na_count = sum(1 for r in results if r["status"] == "N/A")

    if vuln_count >= 3:
        risk_level = "Critical"
    elif vuln_count >= 1:
        risk_level = "High"
    elif warn_count >= 1:
        risk_level = "Medium"
    else:
        risk_level = "Low"

    # 컬럼 너비 통일 (A~E 사용)
    col_widths = {"A": 16, "B": 22, "C": 4, "D": 16, "E": 22}
    for letter, w in col_widths.items():
        ws.column_dimensions[letter].width = w

    # 제목
    ws.merge_cells("A1:E1")
    cell = ws["A1"]
    cell.value = "AWVS 보안 진단 보고서"
    cell.font = FONT_TITLE
    cell.alignment = ALIGN_CENTER
    cell.fill = FILL_DARK_HEADER
    cell.font = Font(name="Arial", bold=True, size=16, color="FFFFFF")
    for c in range(1, 6):
        ws.cell(row=1, column=c).fill = FILL_DARK_HEADER
        ws.cell(row=1, column=c).border = THIN_BORDER
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:E2")
    ws["A2"].value = f"생성: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | 스캔: {scan_data.get('scan_time', '-')}"
    ws["A2"].font = FONT_SMALL
    ws["A2"].alignment = ALIGN_CENTER

    # ── 좌측: 서버 정보 / 우측: 통계 (2열 레이아웃) ──
    row = 4
    # 좌측 헤더
    _section_header(ws, row, 1, 2, "서버 정보", FILL_LIGHT_GRAY)
    # 우측 헤더
    _section_header(ws, row, 4, 5, "진단 결과 통계", FILL_LIGHT_GRAY)

    # 좌측 데이터
    info_data = [
        ("호스트명", scan_data.get("server", "-")),
        ("IP 주소", scan_data.get("server_ip", "-")),
        ("점검 항목 수", f"{len(results)}개"),
    ]
    for i, (label, value) in enumerate(info_data, row + 1):
        _apply_cell(ws, i, 1, label, font=FONT_LABEL)
        _apply_cell(ws, i, 2, value)

    # 우측 데이터
    stat_data = [
        ("종합 위험도", risk_level, Font(name="Arial", bold=True, size=13, color=RISK_COLORS.get(risk_level, "000000")), None),
        ("취약", f"{vuln_count}건", STATUS_FONT_MAP["취약"], FILL_VULN),
        ("주의", f"{warn_count}건", STATUS_FONT_MAP["주의"], FILL_WARN),
        ("양호", f"{safe_count}건", STATUS_FONT_MAP["양호"], FILL_SAFE),
    ]
    for i, (label, value, font, fill) in enumerate(stat_data, row + 1):
        _apply_cell(ws, i, 4, label, font=FONT_LABEL)
        _apply_cell(ws, i, 5, value, font=font, fill=fill, alignment=ALIGN_CENTER)

    # ── 전체 점검 결과 요약 테이블 ──
    table_start = row + len(info_data) + 2
    _section_header(ws, table_start, 1, 5, "전체 점검 결과", FILL_HEADER_BLUE)
    ws.cell(row=table_start, column=1).font = Font(name="Arial", bold=True, size=11, color="FFFFFF")

    table_headers = ["모듈코드", "항목명", "", "상태", "ATT&CK Technique"]
    header_row = table_start + 1
    header_widths_map = {1: "A", 2: "B", 4: "D", 5: "E"}
    for col_idx, header in enumerate(table_headers, 1):
        if col_idx == 3:
            continue  # C열은 간격용
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    # C열 헤더도 색 채우기
    ws.cell(row=header_row, column=3).fill = FILL_HEADER_BLUE
    ws.cell(row=header_row, column=3).border = THIN_BORDER

    for i, r in enumerate(results, header_row + 1):
        status = r.get("status", "")
        ai = r.get("ai_analysis", {})
        technique = ai.get("mitre_technique", "-")

        _apply_cell(ws, i, 1, r.get("module", ""), alignment=ALIGN_CENTER)
        # B+C 병합해서 항목명에 더 넓은 공간
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        _apply_cell(ws, i, 2, r.get("title", ""))
        ws.cell(row=i, column=3).border = THIN_BORDER

        status_cell = _apply_cell(ws, i, 4, status, alignment=ALIGN_CENTER)
        if status in STATUS_FILL_MAP:
            status_cell.fill = STATUS_FILL_MAP[status]
            status_cell.font = STATUS_FONT_MAP[status]

        _apply_cell(ws, i, 5, technique, font=FONT_SMALL)

    # 인쇄 영역 설정
    last_row = header_row + len(results)
    ws.print_area = f"A1:E{last_row}"


# ── Sheet 2: 진단 결과 상세 ──────────────────────────────
def _build_results_sheet(wb, results):
    ws = wb.create_sheet("진단 결과 상세")

    headers = ["모듈코드", "카테고리", "항목명", "점검 대상", "상태", "증거", "판단 근거", "권고사항"]
    widths = [12, 10, 25, 30, 8, 50, 40, 40]
    _apply_header_row(ws, 1, headers, FILL_HEADER_ORANGE, widths)

    for row_idx, r in enumerate(results, 2):
        values = [
            r.get("module", ""),
            r.get("category", "-"),
            r.get("title", ""),
            r.get("target", ""),
            r.get("status", ""),
            r.get("evidence", "").replace(" | ", "\n"),
            r.get("reason", "").replace(" | ", "\n"),
            r.get("recommendation", "").replace(" / ", "\n"),
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_WRAP

        status = r.get("status", "")
        status_cell = ws.cell(row=row_idx, column=5)
        if status in STATUS_FILL_MAP:
            status_cell.fill = STATUS_FILL_MAP[status]
            status_cell.font = STATUS_FONT_MAP[status]
            status_cell.alignment = ALIGN_CENTER

        # 행 높이 자동 조정 (증거 길이 기반)
        evidence_len = len(r.get("evidence", ""))
        if evidence_len > 100:
            ws.row_dimensions[row_idx].height = 60
        elif evidence_len > 50:
            ws.row_dimensions[row_idx].height = 40

    ws.auto_filter.ref = f"A1:H{len(results) + 1}"
    ws.freeze_panes = "A2"


# ── Sheet 3: ATT&CK 매핑 ────────────────────────────────
def _build_attack_sheet(wb, attack_mapping):
    ws = wb.create_sheet("ATT&CK 매핑")

    headers = ["모듈코드", "Tactic", "Technique ID", "Technique Name", "설명", "관련 기법"]
    widths = [12, 18, 15, 35, 50, 25]
    _apply_header_row(ws, 1, headers, FILL_HEADER_PURPLE, widths)

    tactic_colors = {
        "Persistence": "E8D5F5",
        "Collection": "D5E8F5",
        "Discovery": "D5F5E0",
        "Privilege Escalation": "F5E8D5",
        "Credential Access": "F5D5D5",
        "Execution": "F5F0D5",
    }

    sorted_items = sorted(attack_mapping.items(), key=lambda x: x[1].get("tactic", ""))

    for row_idx, (module_id, mapping) in enumerate(sorted_items, 2):
        tactic = mapping.get("tactic", "-")
        technique_id = mapping.get("technique_id", "-")
        related = ", ".join(mapping.get("related", [])) if mapping.get("related") else "-"

        values = [
            module_id,
            tactic,
            technique_id,
            mapping.get("technique_name", "-"),
            mapping.get("description", "-"),
            related,
        ]

        row_fill = tactic_colors.get(tactic)

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_WRAP
            if row_fill:
                cell.fill = PatternFill("solid", fgColor=row_fill)

        # Technique ID에 MITRE 링크
        tech_cell = ws.cell(row=row_idx, column=3)
        if technique_id != "-":
            base_id = technique_id.split(".")[0]
            url = f"https://attack.mitre.org/techniques/{base_id}/"
            if "." in technique_id:
                sub = technique_id.split(".")[1]
                url = f"https://attack.mitre.org/techniques/{base_id}/{sub}/"
            tech_cell.hyperlink = url
            tech_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    ws.auto_filter.ref = f"A1:F{len(attack_mapping) + 1}"
    ws.freeze_panes = "A2"


# ── Sheet 4: 대응 방안 ──────────────────────────────────
def _build_countermeasure_sheet(wb, results):
    ws = wb.create_sheet("대응 방안")

    vuln_items = [r for r in results if r["status"] == "취약"]
    warn_items = [r for r in results if r["status"] == "주의"]

    headers = ["우선순위", "모듈코드", "항목명", "위험 상세", "공격 시나리오", "대응 방안", "ATT&CK Technique"]
    widths = [10, 12, 22, 42, 42, 42, 32]
    _apply_header_row(ws, 1, headers, FILL_HEADER_RED, widths)

    if not vuln_items and not warn_items:
        ws.merge_cells("A2:G2")
        cell = ws["A2"]
        cell.value = "취약/주의 항목이 없습니다. 모든 점검 항목이 양호합니다."
        cell.font = Font(name="Arial", size=11, color="28A745")
        cell.alignment = ALIGN_CENTER
        return

    current_row = 2

    # 취약 항목 (P1, P2, P3...)
    for idx, r in enumerate(vuln_items, 1):
        ai = r.get("ai_analysis", {})
        risk_detail = ai.get("risk_detail", "")
        attack_scenario = ai.get("attack_scenario", "")
        countermeasure = ai.get("countermeasure", "")

        # AI 분석이 비어있으면 스캔 결과에서 대체
        if not countermeasure or countermeasure == "-":
            countermeasure = r.get("recommendation", "-").replace(" / ", "\n")
        if not risk_detail or risk_detail == "-" or risk_detail == "현재 안전한 상태입니다.":
            risk_detail = r.get("reason", "-").replace(" | ", "\n")
        if not attack_scenario or attack_scenario in ("-", "해당 없음"):
            attack_scenario = "-"

        _apply_cell(ws, current_row, 1, f"P{idx}", font=Font(name="Arial", bold=True, size=11, color="CC0000"), fill=FILL_VULN, alignment=ALIGN_CENTER)
        _apply_cell(ws, current_row, 2, r.get("module", ""), fill=FILL_VULN, alignment=ALIGN_CENTER)
        _apply_cell(ws, current_row, 3, r.get("title", ""), fill=FILL_VULN)
        _apply_cell(ws, current_row, 4, risk_detail)
        _apply_cell(ws, current_row, 5, attack_scenario)
        _apply_cell(ws, current_row, 6, countermeasure)
        _apply_cell(ws, current_row, 7, ai.get("mitre_technique", "-"), font=FONT_SMALL)

        ws.row_dimensions[current_row].height = 55
        current_row += 1

    # 주의 항목 구분선 + 데이터
    if warn_items:
        current_row += 1  # 빈 행
        _section_header(ws, current_row, 1, 7, "▼ 주의 항목 (Warning)", FILL_WARN)
        ws.cell(row=current_row, column=1).font = Font(name="Arial", bold=True, size=11, color="856404")
        current_row += 1

        for idx, r in enumerate(warn_items, 1):
            ai = r.get("ai_analysis", {})
            risk_detail = ai.get("risk_detail", "")
            countermeasure = ai.get("countermeasure", "")

            # 주의 항목: AI 분석이 기본값이면 스캔 원본 데이터 사용
            if not countermeasure or countermeasure in ("-", "현재 안전한 상태입니다."):
                countermeasure = r.get("recommendation", "-").replace(" / ", "\n")
            if not risk_detail or risk_detail in ("-", "현재 안전한 상태입니다."):
                risk_detail = r.get("reason", "-").replace(" | ", "\n")

            attack_scenario = ai.get("attack_scenario", "")
            if not attack_scenario or attack_scenario in ("-", "해당 없음"):
                attack_scenario = "-"

            _apply_cell(ws, current_row, 1, f"W{idx}", font=Font(name="Arial", bold=True, size=10, color="856404"), fill=FILL_WARN, alignment=ALIGN_CENTER)
            _apply_cell(ws, current_row, 2, r.get("module", ""), fill=FILL_WARN, alignment=ALIGN_CENTER)
            _apply_cell(ws, current_row, 3, r.get("title", ""), fill=FILL_WARN)
            _apply_cell(ws, current_row, 4, risk_detail)
            _apply_cell(ws, current_row, 5, attack_scenario)
            _apply_cell(ws, current_row, 6, countermeasure)
            _apply_cell(ws, current_row, 7, ai.get("mitre_technique", "-"), font=FONT_SMALL)

            ws.row_dimensions[current_row].height = 45
            current_row += 1

    ws.auto_filter.ref = f"A1:G{current_row - 1}"
    ws.freeze_panes = "A2"


# ── 메인 함수 ────────────────────────────────────────────
def generate_xlsx_report(scan_data, attack_mapping=None, output_dir=None):
    if attack_mapping is None:
        mapping_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "prompts", "attack_mapping.json"
        )
        if os.path.exists(mapping_path):
            with open(mapping_path, "r", encoding="utf-8") as f:
                attack_mapping = json.load(f)
        else:
            attack_mapping = {}

    results = scan_data.get("results", [])
    wb = Workbook()

    _build_summary_sheet(wb, scan_data, results)
    _build_results_sheet(wb, results)
    _build_attack_sheet(wb, attack_mapping)
    _build_countermeasure_sheet(wb, results)

    if output_dir is None:
        output_dir = os.path.expanduser("~/awvs-results")
    os.makedirs(output_dir, exist_ok=True)

    filename = f"awvs_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    return filepath


def generate_xlsx_bytes(scan_data, attack_mapping=None):
    from io import BytesIO

    if attack_mapping is None:
        mapping_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "prompts", "attack_mapping.json"
        )
        if os.path.exists(mapping_path):
            with open(mapping_path, "r", encoding="utf-8") as f:
                attack_mapping = json.load(f)
        else:
            attack_mapping = {}

    results = scan_data.get("results", [])
    wb = Workbook()
    _build_summary_sheet(wb, scan_data, results)
    _build_results_sheet(wb, results)
    _build_attack_sheet(wb, attack_mapping)
    _build_countermeasure_sheet(wb, results)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


if __name__ == "__main__":
    sample_path = os.path.expanduser("~/awvs-results")
    json_files = [f for f in os.listdir(sample_path) if f.endswith(".json")] if os.path.exists(sample_path) else []

    if json_files:
        latest = sorted(json_files)[-1]
        with open(os.path.join(sample_path, latest), "r", encoding="utf-8") as f:
            scan_data = json.load(f)
        path = generate_xlsx_report(scan_data)
        print(f"보고서 생성 완료: {path}")
    else:
        print("스캔 결과 JSON이 없습니다. scanner.py를 먼저 실행하세요.")
